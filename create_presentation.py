"""
create_presentation.py
Generates the complete, professional, visually attractive PowerPoint presentation (.pptx)
for the project:
PROJECT TITLE: AI-ASSISTED RECRUITMENT SYSTEM
STUDENT: S. Srisakthi (310624205269)
DEGREE: B.Tech Information Technology
COLLEGE: Easwari Engineering College (Autonomous)
PROJECT TYPE: Full Stack Development (FSD) Mini Project
DATE: October 2026
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = "AI_Assisted_Recruitment_System_Presentation.pptx"

# Color Palette Constants
DARK_NAVY = RGBColor(11, 19, 43)        # #0B132B
SLATE_DARK = RGBColor(28, 37, 65)       # #1C2541
CARD_DARK = RGBColor(30, 41, 59)        # #1E293B
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252)      # #F8FAFC
CARD_BG = RGBColor(255, 255, 255)       # #FFFFFF
TEXT_PRIMARY = RGBColor(15, 23, 42)     # #0F172A
TEXT_MUTED = RGBColor(100, 116, 139)    # #64748B
TEXT_BODY = RGBColor(51, 65, 85)        # #334155
ROYAL_BLUE = RGBColor(37, 99, 235)      # #2563EB
BLUE_DARK = RGBColor(29, 78, 216)       # #1D4ED8
CYAN_ACCENT = RGBColor(14, 165, 233)    # #0EA5E9
CYAN_LIGHT = RGBColor(56, 189, 248)     # #38BDF8
EMERALD = RGBColor(16, 185, 129)        # #10B981
EMERALD_BG = RGBColor(236, 253, 245)    # #ECFDF5
BORDER_COLOR = RGBColor(226, 232, 240)  # #E2E8F0
BORDER_DARK = RGBColor(51, 65, 85)      # #334155
TABLE_HEADER_BG = RGBColor(15, 23, 42)  # #0F172A
TABLE_ROW_ALT = RGBColor(241, 245, 249) # #F1F5F9


def init_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs


def set_slide_background(slide, color):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
    )
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def add_header(slide, title_text, kicker_text=""):
    """Adds a standard executive header to content slides."""
    # Kicker / Category tag
    if kicker_text:
        k_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.40), Inches(11.7), Inches(0.3))
        tf_k = k_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_right = tf_k.margin_top = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = kicker_text.upper()
        p_k.font.size = Pt(9.5)
        p_k.font.bold = True
        p_k.font.color.rgb = ROYAL_BLUE
        p_k.font.name = "Segoe UI"

    # Slide Title
    t_top = Inches(0.65) if kicker_text else Inches(0.5)
    t_box = slide.shapes.add_textbox(Inches(0.8), t_top, Inches(11.7), Inches(0.55))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(22)
    p_t.font.bold = True
    p_t.font.color.rgb = TEXT_PRIMARY
    p_t.font.name = "Segoe UI"

    # Thin accent divider
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.28), Inches(11.733), Inches(0.025)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = ROYAL_BLUE
    line.line.fill.background()


def add_footer(slide, slide_num, total_slides=15):
    """Adds standard footer to content slides."""
    # Divider line
    f_line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.015)
    )
    f_line.fill.solid()
    f_line.fill.fore_color.rgb = BORDER_COLOR
    f_line.line.fill.background()

    # Footer Left
    l_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(4.5), Inches(0.3))
    tf_l = l_box.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = tf_l.margin_right = tf_l.margin_top = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "AI-Assisted Recruitment System  |  Full Stack Development Mini Project"
    p_l.font.size = Pt(8.5)
    p_l.font.color.rgb = TEXT_MUTED
    p_l.font.name = "Segoe UI"

    # Footer Center
    c_box = slide.shapes.add_textbox(Inches(5.0), Inches(6.92), Inches(5.5), Inches(0.3))
    tf_c = c_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = tf_c.margin_right = tf_c.margin_top = tf_c.margin_bottom = 0
    p_c = tf_c.paragraphs[0]
    p_c.alignment = PP_ALIGN.CENTER
    p_c.text = "S. Srisakthi (310624205269)  |  Easwari Engineering College"
    p_c.font.size = Pt(8.5)
    p_c.font.color.rgb = TEXT_MUTED
    p_c.font.name = "Segoe UI"

    # Footer Right
    r_box = slide.shapes.add_textbox(Inches(11.0), Inches(6.92), Inches(1.533), Inches(0.3))
    tf_r = r_box.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = tf_r.margin_right = tf_r.margin_top = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.text = f"Slide {slide_num} / {total_slides}"
    p_r.font.size = Pt(8.5)
    p_r.font.bold = True
    p_r.font.color.rgb = TEXT_PRIMARY
    p_r.font.name = "Segoe UI"


def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR, border_pt=1):
    """Creates a rounded container card for visual grouping."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_pt)
    else:
        card.line.fill.background()
    return card


def add_bullet_points(tf, points, font_size=10, text_color=TEXT_BODY, space_after=6, bold_first_n_words=0):
    """Adds cleanly formatted bullet points to a text frame."""
    for i, pt_text in enumerate(points):
        p = tf.add_paragraph() if i > 0 or len(tf.paragraphs[0].text) > 0 else tf.paragraphs[0]
        p.space_after = Pt(space_after)
        p.space_before = Pt(2)
        p.font.size = Pt(font_size)
        p.font.name = "Segoe UI"
        p.font.color.rgb = text_color

        if bold_first_n_words > 0 and ":" in pt_text:
            parts = pt_text.split(":", 1)
            r1 = p.add_run()
            r1.text = "• " + parts[0] + ":"
            r1.font.bold = True
            r1.font.color.rgb = TEXT_PRIMARY
            r1.font.size = Pt(font_size)
            r1.font.name = "Segoe UI"

            r2 = p.add_run()
            r2.text = parts[1]
            r2.font.bold = False
            r2.font.color.rgb = text_color
            r2.font.size = Pt(font_size)
            r2.font.name = "Segoe UI"
        else:
            p.text = "• " + pt_text


# ==============================================================================
# SLIDE BUILDERS
# ==============================================================================

def build_slide_1(prs):
    """SLIDE 1 — TITLE SLIDE (Dark Navy Theme)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, DARK_NAVY)

    # Decorative background accent elements
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.12), Inches(5.8)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = CYAN_ACCENT
    accent_bar.line.fill.background()

    # Kicker Badge
    k_badge = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.9), Inches(4.5), Inches(0.38)
    )
    k_badge.fill.solid()
    k_badge.fill.fore_color.rgb = BLUE_DARK
    k_badge.line.fill.background()
    tf_kb = k_badge.text_frame
    p_kb = tf_kb.paragraphs[0]
    p_kb.text = "FULL STACK DEVELOPMENT (FSD) MINI PROJECT"
    p_kb.font.size = Pt(10)
    p_kb.font.bold = True
    p_kb.font.color.rgb = WHITE
    p_kb.font.name = "Segoe UI"
    p_kb.alignment = PP_ALIGN.CENTER

    # Project Title
    t_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.4), Inches(7.5), Inches(1.4))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "AI-ASSISTED\nRECRUITMENT SYSTEM"
    p_t.font.size = Pt(36)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE
    p_t.font.name = "Segoe UI"

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.0), Inches(7.5), Inches(0.9))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "An Intelligent MERN Stack Platform with Modular NLP Resume Ingestion, Explainable Multi-Factor Candidate Ranking & 6-Stage Kanban Pipeline"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = CYAN_LIGHT
    p_sub.font.name = "Segoe UI"

    # Technology Pill Badges
    tech_tags = ["MongoDB", "Express.js", "React 19", "Node.js", "Compromise NLP", "Natural TF-IDF"]
    start_x = Inches(1.2)
    for tag in tech_tags:
        pill = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, start_x, Inches(4.1), Inches(1.15), Inches(0.35)
        )
        pill.fill.solid()
        pill.fill.fore_color.rgb = SLATE_DARK
        pill.line.color.rgb = CYAN_ACCENT
        pill.line.width = Pt(1)
        tf_p = pill.text_frame
        p_p = tf_p.paragraphs[0]
        p_p.text = tag
        p_p.font.size = Pt(8.5)
        p_p.font.bold = True
        p_p.font.color.rgb = WHITE
        p_p.font.name = "Segoe UI"
        p_p.alignment = PP_ALIGN.CENTER
        start_x += Inches(1.23)

    # Student Credentials Card (Right side)
    cred_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.2), Inches(3.6), Inches(5.1)
    )
    cred_card.fill.solid()
    cred_card.fill.fore_color.rgb = CARD_DARK
    cred_card.line.color.rgb = BORDER_DARK
    cred_card.line.width = Pt(1.5)

    tf_cred = cred_card.text_frame
    tf_cred.word_wrap = True
    tf_cred.margin_left = Inches(0.25)
    tf_cred.margin_right = Inches(0.25)
    tf_cred.margin_top = Inches(0.25)

    # Header inside card
    p_ch = tf_cred.paragraphs[0]
    p_ch.text = "CANDIDATE CREDENTIALS"
    p_ch.font.size = Pt(10.5)
    p_ch.font.bold = True
    p_ch.font.color.rgb = CYAN_LIGHT
    p_ch.font.name = "Segoe UI"
    p_ch.space_after = Pt(14)

    fields = [
        ("Student Name", "S. Srisakthi"),
        ("Register Number", "310624205269"),
        ("Degree & Branch", "B.Tech. Information Technology"),
        ("Department", "Information Technology (IT)"),
        ("Institution", "Easwari Engineering College\n(Autonomous), Chennai"),
        ("Project Category", "Full Stack Development (FSD)"),
        ("Academic Date", "October 2026"),
    ]

    for label, val in fields:
        p_lbl = tf_cred.add_paragraph()
        p_lbl.text = label.upper()
        p_lbl.font.size = Pt(8)
        p_lbl.font.color.rgb = TEXT_MUTED
        p_lbl.font.name = "Segoe UI"
        p_lbl.space_after = Pt(1)

        p_val = tf_cred.add_paragraph()
        p_val.text = val
        p_val.font.size = Pt(10.5)
        p_val.font.bold = True
        p_val.font.color.rgb = WHITE
        p_val.font.name = "Segoe UI"
        p_val.space_after = Pt(8)


def build_slide_2(prs):
    """SLIDE 2 — INTRODUCTION"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Introduction to AI-Assisted Recruitment", "Overview & Context")
    add_footer(slide, 2)

    card_w = Inches(3.7)
    card_h = Inches(4.1)
    card_y = Inches(1.55)

    # Card 1: Modern Recruitment Landscape
    create_card(slide, Inches(0.8), card_y, card_w, card_h)
    tb1 = slide.shapes.add_textbox(Inches(0.95), Inches(1.7), Inches(3.4), Inches(3.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "1. Modern Hiring Landscape"
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = ROYAL_BLUE
    p1.space_after = Pt(10)
    pts1 = [
        "High Application Volumes: Digital career boards produce hundreds of unstructured resumes for every open position.",
        "Cognitive Fatigue: Recruiters spend 3–5 minutes manually scanning each resume, leading to burnout and subjective inconsistencies.",
        "Delayed Time-to-Hire: Traditional manual triage bottlenecks modern hiring cycles and increases candidate churn.",
        "Communication Gaps: Disconnected workflows leave applicants uninformed and recruiters overburdened."
    ]
    add_bullet_points(tf1, pts1, font_size=9.5, bold_first_n_words=1)

    # Card 2: Purpose of Proposed System
    create_card(slide, Inches(4.8), card_y, card_w, card_h)
    tb2 = slide.shapes.add_textbox(Inches(4.95), Inches(1.7), Inches(3.4), Inches(3.8))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    p2 = tf2.paragraphs[0]
    p2.text = "2. Purpose of the System"
    p2.font.size = Pt(13)
    p2.font.bold = True
    p2.font.color.rgb = ROYAL_BLUE
    p2.space_after = Pt(10)
    pts2 = [
        "Automated Resume Ingestion: Eliminates manual data entry by extracting structured candidate profiles directly from PDF, DOCX, and TXT resumes.",
        "Explainable Candidate Matching: Computes multi-factor compatibility scores across skills, experience, education, and domain text fit.",
        "Assistive Decision Support: Provides objective recommendations (Strong, Good, Potential, Low) without replacing human recruiters.",
        "Interactive 6-Stage Pipeline: Provides an end-to-end Kanban board to manage candidates from application to interview scheduling."
    ]
    add_bullet_points(tf2, pts2, font_size=9.5, bold_first_n_words=1)

    # Card 3: Role of Full-Stack Development
    create_card(slide, Inches(8.8), card_y, card_w, card_h)
    tb3 = slide.shapes.add_textbox(Inches(8.95), Inches(1.7), Inches(3.4), Inches(3.8))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    p3 = tf3.paragraphs[0]
    p3.text = "3. Full-Stack MERN Architecture"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = ROYAL_BLUE
    p3.space_after = Pt(10)
    pts3 = [
        "React 19 Frontend: Responsive Single Page Application with dynamic state, real-time filtering, and role-based views.",
        "Node.js & Express REST API: Asynchronous backend controllers enforcing JWT authentication and business logic.",
        "MongoDB Document Store: Schema-flexible database indexing nested candidate profiles, jobs, and application histories.",
        "Seamless System Integration: Connects client interactions, NLP parsing services, and database persistence in one cohesive stack."
    ]
    add_bullet_points(tf3, pts3, font_size=9.5, bold_first_n_words=1)

    # Bottom highlight banner
    ban = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.85), Inches(11.733), Inches(0.8)
    )
    ban.fill.solid()
    ban.fill.fore_color.rgb = RGBColor(238, 242, 255)
    ban.line.color.rgb = ROYAL_BLUE
    ban.line.width = Pt(1)
    tf_b = ban.text_frame
    tf_b.word_wrap = True
    p_b = tf_b.paragraphs[0]
    p_b.text = "Core Value: AI Recruit Pro bridges the gap between massive applicant volumes and recruiter productivity through modular natural language processing, transparent scoring, and structured pipeline management."
    p_b.font.size = Pt(10)
    p_b.font.color.rgb = BLUE_DARK
    p_b.font.name = "Segoe UI"
    p_b.alignment = PP_ALIGN.CENTER


def build_slide_3(prs):
    """SLIDE 3 — PROBLEM STATEMENT"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Recruitment Challenges Addressed", "Problem Statement & Industry Context")
    add_footer(slide, 3)

    # 4 Problem Cards in 2x2 grid
    card_w = Inches(5.7)
    card_h = Inches(2.25)
    xs = [Inches(0.8), Inches(6.8)]
    ys = [Inches(1.5), Inches(3.95)]

    problems = [
        (
            "1. Manual & Labor-Intensive Resume Screening",
            "Recruiters manually download and review diverse resume formats (PDF, DOCX, TXT) line-by-line.",
            [
                "Requires 3–5 minutes per candidate; screening 200 applicants takes over 12 recruiter hours.",
                "High cognitive fatigue leads to overlooked talent, inconsistent assessments, and hiring delays.",
                "Unstructured text makes comparative candidate evaluation difficult and error-prone."
            ]
        ),
        (
            "2. Subjective & Inconsistent Candidate Evaluation",
            "Screening decisions often rely on subjective keyword searches without standard evaluation criteria.",
            [
                "Recruiters struggle to simultaneously balance skills, experience years, and degree qualifications.",
                "Varying evaluation standards across team members lead to biased or unfair shortlisting.",
                "Absence of transparent, mathematical scoring criteria prevents objective candidate ranking."
            ]
        ),
        (
            "3. Disconnected Candidate Data Management",
            "Candidate profiles, resumes, and interview feedback are scattered across disparate platforms.",
            [
                "Resumes stored in email threads, local desktop folders, and isolated spreadsheets.",
                "Risk of duplicate applications, misplaced candidate data, and zero centralized talent records.",
                "Absence of an integrated audit trail tracking who reviewed, advanced, or rejected an applicant."
            ]
        ),
        (
            "4. Opaque Pipeline & Disjointed Scheduling",
            "Traditional workflows fail to provide pipeline visibility for recruiters and candidates alike.",
            [
                "Candidates experience a 'black hole' with zero feedback on their current hiring stage.",
                "Lack of an interactive visual Kanban board to track applicants across hiring phases.",
                "Interview coordination requires external back-and-forth emails, causing high candidate drop-off."
            ]
        ),
    ]

    idx = 0
    for r in range(2):
        for c in range(2):
            title, desc, bullets = problems[idx]
            create_card(slide, xs[c], ys[r], card_w, card_h)

            tb = slide.shapes.add_textbox(xs[c] + Inches(0.18), ys[r] + Inches(0.12), card_w - Inches(0.36), card_h - Inches(0.24))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(11.5)
            p.font.bold = True
            p.font.color.rgb = ROYAL_BLUE
            p.space_after = Pt(3)

            p_desc = tf.add_paragraph()
            p_desc.text = desc
            p_desc.font.size = Pt(9.5)
            p_desc.font.italic = True
            p_desc.font.color.rgb = TEXT_MUTED
            p_desc.space_after = Pt(5)

            add_bullet_points(tf, bullets, font_size=9, space_after=3)
            idx += 1

    # Realistic Scope Callout
    callout = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.35), Inches(11.733), Inches(0.42)
    )
    callout.fill.solid()
    callout.fill.fore_color.rgb = RGBColor(254, 243, 199)
    callout.line.color.rgb = RGBColor(245, 158, 11)
    callout.line.width = Pt(1)
    tf_co = callout.text_frame
    p_co = tf_co.paragraphs[0]
    p_co.text = "Realistic Boundary: AI Recruit Pro serves as an assistive decision-support platform to accelerate triage and ranking — final hiring decisions remain with human recruiters."
    p_co.font.size = Pt(9)
    p_co.font.bold = True
    p_co.font.color.rgb = RGBColor(146, 64, 14)
    p_co.alignment = PP_ALIGN.CENTER


def build_slide_4(prs):
    """SLIDE 4 — PROJECT OBJECTIVES"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Core Project Objectives", "Goals & Functional Deliverables")
    add_footer(slide, 4)

    card_w = Inches(5.7)
    card_h = Inches(2.45)
    xs = [Inches(0.8), Inches(6.8)]
    ys = [Inches(1.5), Inches(4.15)]

    objectives = [
        (
            "01",
            "Modular Resume Ingestion & Parsing",
            "Develop a robust text-processing service that extracts candidate data from multiple file formats.",
            [
                "Multi-Format Ingestion: Ingest PDF, DOCX, and TXT resumes via Multer file streaming.",
                "Entity Recognition: Extract applicant contact details (email, phone, links) using regex & Compromise NLP.",
                "Taxonomy Extraction: Identify skills against a curated 500+ technical skill taxonomy.",
                "Profile Normalization: Calculate total years of experience and classify highest degree."
            ]
        ),
        (
            "02",
            "Explainable Multi-Factor Candidate Ranking",
            "Implement an objective, mathematical scoring engine to evaluate applicant-job compatibility.",
            [
                "Multi-Criteria Weighting: 50% Skills Overlap + 25% Experience + 15% Education + 10% Semantic Fit.",
                "Alias Mapping: Match synonym skill variations (e.g., K8s -> Kubernetes, React.js -> React).",
                "Transparent Breakdown: Provide recruiters with itemized sub-scores and missing skill badges.",
                "Recommendation Tiers: Automatically categorize candidates into Strong, Good, Potential, or Low Match."
            ]
        ),
        (
            "03",
            "Interactive 6-Stage Kanban Hiring Pipeline",
            "Build an intuitive, interactive pipeline to advance applicants smoothly through hiring rounds.",
            [
                "Structured Stages: Applied ➔ Screened ➔ Shortlisted ➔ Interview Scheduled ➔ Selected / Rejected.",
                "Audit History: Maintain complete transition logs with timestamps, stage notes, and updater identity.",
                "Recruiter Oversight: Enable 1-click stage advancement directly from applicant cards.",
                "Data Integrity: Prevent duplicate applications via unique compound MongoDB indexing."
            ]
        ),
        (
            "04",
            "Interview Scheduling & Candidate Portal",
            "Unify interview coordination and provide transparent status feedback for job applicants.",
            [
                "Integrated Scheduling: Configure interview date, type (Technical/HR), Google Meet links, and notes.",
                "Candidate Dashboard: Display real-time application stage progress bars and match insights.",
                "Meeting Access: Provide direct meeting links and interviewer details on the candidate interface.",
                "Role-Based Security: Enforce JWT authentication with strict role authorization guards."
            ]
        ),
    ]

    idx = 0
    for r in range(2):
        for c in range(2):
            num, title, subtitle, bullets = objectives[idx]
            create_card(slide, xs[c], ys[r], card_w, card_h)

            # Number badge
            nb = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, xs[c] + Inches(0.18), ys[r] + Inches(0.15), Inches(0.55), Inches(0.35)
            )
            nb.fill.solid()
            nb.fill.fore_color.rgb = ROYAL_BLUE
            nb.line.fill.background()
            tf_nb = nb.text_frame
            p_nb = tf_nb.paragraphs[0]
            p_nb.text = num
            p_nb.font.size = Pt(11)
            p_nb.font.bold = True
            p_nb.font.color.rgb = WHITE
            p_nb.alignment = PP_ALIGN.CENTER

            # Title
            t_box = slide.shapes.add_textbox(xs[c] + Inches(0.85), ys[r] + Inches(0.12), card_w - Inches(1.0), Inches(0.4))
            tf_t = t_box.text_frame
            tf_t.word_wrap = True
            p_t = tf_t.paragraphs[0]
            p_t.text = title
            p_t.font.size = Pt(12)
            p_t.font.bold = True
            p_t.font.color.rgb = TEXT_PRIMARY

            # Subtitle
            s_box = slide.shapes.add_textbox(xs[c] + Inches(0.18), ys[r] + Inches(0.55), card_w - Inches(0.36), Inches(0.3))
            tf_s = s_box.text_frame
            tf_s.word_wrap = True
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle
            p_s.font.size = Pt(9)
            p_s.font.italic = True
            p_s.font.color.rgb = TEXT_MUTED

            # Bullets
            b_box = slide.shapes.add_textbox(xs[c] + Inches(0.18), ys[r] + Inches(0.85), card_w - Inches(0.36), card_h - Inches(0.95))
            tf_b = b_box.text_frame
            tf_b.word_wrap = True
            add_bullet_points(tf_b, bullets, font_size=8.5, space_after=3)

            idx += 1


def build_slide_5(prs):
    """SLIDE 5 — EXISTING SYSTEM AND PROPOSED SYSTEM"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Existing System vs. Proposed System", "Comparative Analysis & System Benefits")
    add_footer(slide, 5)

    # Comparison Table
    table_shape = slide.shapes.add_table(7, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.1))
    table = table_shape.table

    table.columns[0].width = Inches(2.1)
    table.columns[1].width = Inches(3.3)
    table.columns[2].width = Inches(3.8)
    table.columns[3].width = Inches(2.533)

    headers = [
        "Evaluation Dimension",
        "Traditional / Existing Approach",
        "Proposed System (AI Recruit Pro)",
        "Key Operational Benefit"
    ]

    for col_idx, h_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEADER_BG
        cell.text_frame.word_wrap = True
        p = cell.text_frame.paragraphs[0]
        p.text = h_text
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.font.name = "Segoe UI"
        p.alignment = PP_ALIGN.LEFT

    rows_data = [
        (
            "Resume Screening",
            "Manual line-by-line reading of unstructured documents; 3–5 min per resume.",
            "Automated modular NLP parsing (pdf-parse, mammoth, compromise) extracting contact, skills & exp.",
            "Immediate profile extraction; eliminates manual data entry."
        ),
        (
            "Candidate Evaluation",
            "Subjective, unstandardized keyword scanning; inconsistent evaluator standards.",
            "Explainable 4-Factor Scoring Formula (50% Skills, 25% Exp, 15% Edu, 10% TF-IDF Fit).",
            "Standardized, unbiased 0–100% compatibility score."
        ),
        (
            "Pipeline Management",
            "Fragmented spreadsheets, email threads, and untracked desktop files.",
            "Interactive 6-Stage Kanban Board with stage transition logs and audit notes.",
            "Complete hiring visibility and historical tracking."
        ),
        (
            "Interview Scheduling",
            "Manual calendar coordination via back-and-forth external emails.",
            "Integrated scheduler supporting meeting links (Google Meet), round types & feedback notes.",
            "Single-platform interview setup without leaving workflow."
        ),
        (
            "Applicant Transparency",
            "Zero status feedback ('black hole' experience); high candidate dissatisfaction.",
            "Dedicated candidate portal with real-time stage progress bars and match insights.",
            "Transparent candidate experience with instant status updates."
        ),
        (
            "Data Storage & Integrity",
            "Siloed file attachments and unindexed spreadsheets risking data loss.",
            "Centralized MongoDB document store with indexed relational schemas and unique constraints.",
            "Prevents duplicate applications; fast and persistent storage."
        ),
    ]

    for row_idx, row_items in enumerate(rows_data, start=1):
        for col_idx, cell_text in enumerate(row_items):
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TABLE_ROW_ALT if row_idx % 2 == 1 else WHITE
            cell.text_frame.word_wrap = True
            cell.text_frame.margin_left = Inches(0.1)
            cell.text_frame.margin_right = Inches(0.1)
            cell.text_frame.margin_top = Inches(0.08)
            cell.text_frame.margin_bottom = Inches(0.08)
            p = cell.text_frame.paragraphs[0]
            p.text = cell_text
            p.font.size = Pt(8.5)
            p.font.name = "Segoe UI"
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = ROYAL_BLUE
            elif col_idx == 2:
                p.font.color.rgb = TEXT_PRIMARY
            else:
                p.font.color.rgb = TEXT_BODY


def build_slide_6(prs):
    """SLIDE 6 — SYSTEM ARCHITECTURE"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "System Architecture & Request Flow", "Technical Design & Layer Interactions")
    add_footer(slide, 6)

    # Embed Architecture Diagram Image (Left / Center)
    arch_img_path = "diagrams/system_architecture.png"
    if os.path.exists(arch_img_path):
        create_card(slide, Inches(0.8), Inches(1.48), Inches(6.8), Inches(5.15))
        slide.shapes.add_picture(arch_img_path, Inches(0.9), Inches(1.58), width=Inches(6.6))
    else:
        # Fallback card if diagram missing
        create_card(slide, Inches(0.8), Inches(1.48), Inches(6.8), Inches(5.15))

    # Architecture Layer Breakdown Cards (Right Column)
    layer_w = Inches(4.7)
    layer_h = Inches(1.2)
    start_y = Inches(1.48)
    gap = Inches(1.32)

    layers = [
        (
            "Layer 1: Presentation (React 19 Frontend)",
            ROYAL_BLUE,
            [
                "Vite-powered Single Page Application (SPA) with responsive dark theme.",
                "Tailwind CSS v4 layouts; Axios client with JWT Bearer interceptor.",
                "Recruiter dashboards, interactive Kanban boards, and candidate portals."
            ]
        ),
        (
            "Layer 2: Application API (Node.js & Express)",
            BLUE_DARK,
            [
                "Modular REST API controllers handling Auth, Jobs, Resumes, and Pipeline.",
                "Role-Based Access Control (RBAC) middleware verifying JWT tokens.",
                "Multer streaming engine handling multipart PDF/DOCX file uploads."
            ]
        ),
        (
            "Layer 3: Modular AI & NLP Services",
            EMERALD,
            [
                "resumeParser.js: Extracts text via pdf-parse/mammoth; parses entities & 500+ skills.",
                "matchingEngine.js: Evaluates 4-factor compatibility formula with Natural TF-IDF.",
                "Decoupled contract ready for plug-and-play LLM (Gemini / OpenAI) replacement."
            ]
        ),
        (
            "Layer 4: Persistence Layer (MongoDB)",
            RGBColor(13, 148, 136),
            [
                "Document collections: Users, Jobs, CandidateProfiles, and Applications.",
                "Mongoose ODM schema validation, typed models, and cross-reference population.",
                "Compound unique indexing preventing duplicate candidate submissions."
            ]
        ),
    ]

    for i, (title, color, bullets) in enumerate(layers):
        y = start_y + (i * gap)
        create_card(slide, Inches(7.833), y, layer_w, layer_h)

        tb = slide.shapes.add_textbox(Inches(7.98), y + Inches(0.08), layer_w - Inches(0.3), layer_h - Inches(0.16))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(2)

        add_bullet_points(tf, bullets, font_size=8, space_after=1)


def build_slide_7(prs):
    """SLIDE 7 — TECHNOLOGIES USED"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Technology Stack & Architectural Rationale", "Core Technologies & Justifications")
    add_footer(slide, 7)

    card_w = Inches(5.7)
    card_h = Inches(2.45)
    xs = [Inches(0.8), Inches(6.8)]
    ys = [Inches(1.5), Inches(4.15)]

    stacks = [
        (
            "Frontend Architecture (Client)",
            ROYAL_BLUE,
            [
                "React 19 (Vite): Provides component modularity, instant HMR, and responsive virtual DOM reconciliation.",
                "Tailwind CSS v4: Delivers utility-first styling for dark-mode modern UI with zero runtime stylesheet overhead.",
                "React Router v7: Enables client-side declarative routing and protected route access guards.",
                "Axios HTTP Client: Intercepts requests to inject JWT Bearer tokens and handles REST communication.",
                "Lucide React & Canvas-Confetti: Delivers consistent iconography and celebration feedback on success."
            ]
        ),
        (
            "Backend Architecture (Server API)",
            BLUE_DARK,
            [
                "Node.js Runtime: Asynchronous event-driven architecture supporting concurrent HTTP requests efficiently.",
                "Express.js Framework: Lightweight routing layer structuring controllers, services, and middlewares cleanly.",
                "JWT (jsonwebtoken): Stateless Bearer token authorization securing recruiter and candidate routes.",
                "BcryptJS: Performs 10-round salted one-way hashing for secure user password storage.",
                "Multer Middleware: Manages streaming multipart/form-data resume file uploads with size and MIME validation."
            ]
        ),
        (
            "Database & Data Modeling",
            RGBColor(13, 148, 136),
            [
                "MongoDB: High-performance NoSQL document database suited for nested, polymorphic resume JSON profiles.",
                "Mongoose ODM: Strongly typed schema enforcement, data validation, and automated timestamp generation.",
                "Model Population: Efficiently resolves cross-document references (Applications ➔ Jobs ➔ CandidateProfiles).",
                "Indexing Strategy: Compound unique index on { job: 1, candidate: 1 } prevents redundant job applications."
            ]
        ),
        (
            "AI / NLP & Quality Assurance",
            EMERALD,
            [
                "Compromise NLP: Performs lightweight English entity recognition to extract candidate names and organizations.",
                "Natural Library (TF-IDF): Evaluates statistical term frequency between resume content and job descriptions.",
                "pdf-parse & mammoth: Low-level binary buffer parsers extracting clean text from PDF and Word documents.",
                "Curated 500+ Skill Taxonomy: Categorizes technical skills across Web, Cloud, DevOps, AI/ML, and Mobile.",
                "Node.js Native Test Suite: Integration test suite (api.test.js) validating 12 end-to-end API and NLP assertions."
            ]
        ),
    ]

    idx = 0
    for r in range(2):
        for c in range(2):
            title, color, bullets = stacks[idx]
            create_card(slide, xs[c], ys[r], card_w, card_h)

            tb = slide.shapes.add_textbox(xs[c] + Inches(0.2), ys[r] + Inches(0.12), card_w - Inches(0.4), card_h - Inches(0.24))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = color
            p.space_after = Pt(6)

            add_bullet_points(tf, bullets, font_size=8.5, space_after=3, bold_first_n_words=1)
            idx += 1


def build_slide_8(prs):
    """SLIDE 8 — SYSTEM MODULES"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Core System Modules & Functional Capabilities", "Modular System Breakdown")
    add_footer(slide, 8)

    card_w = Inches(3.7)
    card_h = Inches(2.45)
    xs = [Inches(0.8), Inches(4.8), Inches(8.8)]
    ys = [Inches(1.5), Inches(4.15)]

    modules = [
        (
            "1. Authentication & RBAC",
            "authController.js / auth.js",
            ROYAL_BLUE,
            [
                "Role-based authentication distinguishing Recruiters and Candidates.",
                "JWT token signing and verification middleware protecting private routes.",
                "Built-in 1-click demo switcher enabling rapid evaluation between roles.",
                "Bcrypt salted password encryption (10 rounds)."
            ]
        ),
        (
            "2. Job Postings Management",
            "jobController.js / Job.js",
            BLUE_DARK,
            [
                "Complete CRUD operations for recruiter job listings.",
                "Tag-based required and optional skill array management.",
                "Experience thresholds (min/max years) and degree requirement rules.",
                "Live applicant counter automatically updated per job."
            ]
        ),
        (
            "3. Modular Resume NLP Parser",
            "resumeParser.js / resumeController.js",
            EMERALD,
            [
                "Ingests PDF, DOCX, and TXT resumes via file upload or raw paste.",
                "Regex & Compromise NLP for candidate contact info and names.",
                "500+ curated tech skills taxonomy lookup and alias mapping.",
                "Calculates total experience duration and highest education degree."
            ]
        ),
        (
            "4. Multi-Factor Matching Engine",
            "matchingEngine.js / matchingController.js",
            RGBColor(147, 51, 234),
            [
                "Calculates 0–100% compatibility score via 4-factor weighted formula.",
                "Evaluates skills overlap (50%), experience (25%), and education (15%).",
                "Computes Natural TF-IDF semantic relevance (10%).",
                "Categorizes into Strong, Good, Potential, or Low Match tiers."
            ]
        ),
        (
            "5. 6-Stage Kanban Pipeline",
            "applicationController.js / Application.js",
            RGBColor(217, 119, 6),
            [
                "Interactive workflow: Applied ➔ Screened ➔ Shortlisted ➔ Interview Scheduled ➔ Selected / Rejected.",
                "Audit logs tracking stage transitions, timestamps, and updater notes.",
                "1-click candidate advancement directly on recruiter board.",
                "Filters applicants by job, stage, and match score."
            ]
        ),
        (
            "6. Interview Scheduler & Portal",
            "applicationController.js / statsController.js",
            RGBColor(13, 148, 136),
            [
                "Built-in interview scheduling with Google Meet links and round types.",
                "Candidate dashboard displaying real-time application status progress.",
                "Direct video call access and recruiter feedback review.",
                "Recruiter overview metrics (active jobs, applicants, avg match score)."
            ]
        ),
    ]

    idx = 0
    for r in range(2):
        for c in range(3):
            title, file_ref, color, bullets = modules[idx]
            create_card(slide, xs[c], ys[r], card_w, card_h)

            tb = slide.shapes.add_textbox(xs[c] + Inches(0.18), ys[r] + Inches(0.12), card_w - Inches(0.36), card_h - Inches(0.24))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = color

            p_ref = tf.add_paragraph()
            p_ref.text = file_ref
            p_ref.font.size = Pt(8)
            p_ref.font.italic = True
            p_ref.font.color.rgb = TEXT_MUTED
            p_ref.space_after = Pt(4)

            add_bullet_points(tf, bullets, font_size=8, space_after=2)
            idx += 1


def build_slide_9(prs):
    """SLIDE 9 — WORKFLOW"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "End-to-End Recruitment Workflow", "Operational Lifecycle & Sequence")
    add_footer(slide, 9)

    # 7 Step Horizontal Flow Cards
    steps = [
        ("Step 1", "Job Creation", "Recruiter posts job specifying required skills, experience bounds & education.", ROYAL_BLUE),
        ("Step 2", "Resume Upload", "Candidate uploads resume (PDF/DOCX/TXT) or pastes raw resume text.", BLUE_DARK),
        ("Step 3", "NLP Parsing", "Modular parser extracts contact data, 500+ skills, exp duration & degree.", EMERALD),
        ("Step 4", "Job Application", "Candidate applies; system immediately executes multi-factor evaluation.", RGBColor(147, 51, 234)),
        ("Step 5", "Match Ranking", "Recruiter inspects ranked candidate list sorted descending by match score.", RGBColor(217, 119, 6)),
        ("Step 6", "Kanban Triage", "Recruiter advances candidate through 6 Kanban stages with audit notes.", RGBColor(13, 148, 136)),
        ("Step 7", "Interview Setup", "Recruiter schedules interview with Google Meet link; candidate tracks live.", RGBColor(225, 29, 72)),
    ]

    step_w = Inches(1.52)
    step_h = Inches(3.2)
    start_x = Inches(0.8)
    gap = Inches(1.7)

    for i, (s_num, s_title, s_desc, color) in enumerate(steps):
        x = start_x + (i * gap)
        create_card(slide, x, Inches(1.6), step_w, step_h)

        # Header badge
        badge = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, x + Inches(0.12), Inches(1.75), step_w - Inches(0.24), Inches(0.35)
        )
        badge.fill.solid()
        badge.fill.fore_color.rgb = color
        badge.line.fill.background()
        tf_bg = badge.text_frame
        p_bg = tf_bg.paragraphs[0]
        p_bg.text = s_num
        p_bg.font.size = Pt(9.5)
        p_bg.font.bold = True
        p_bg.font.color.rgb = WHITE
        p_bg.alignment = PP_ALIGN.CENTER

        # Title
        t_box = slide.shapes.add_textbox(x + Inches(0.08), Inches(2.2), step_w - Inches(0.16), Inches(0.6))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = s_title
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_PRIMARY
        p_t.alignment = PP_ALIGN.CENTER

        # Description
        d_box = slide.shapes.add_textbox(x + Inches(0.08), Inches(2.85), step_w - Inches(0.16), Inches(1.8))
        tf_d = d_box.text_frame
        tf_d.word_wrap = True
        p_d = tf_d.paragraphs[0]
        p_d.text = s_desc
        p_d.font.size = Pt(8.5)
        p_d.font.color.rgb = TEXT_BODY
        p_d.alignment = PP_ALIGN.CENTER

        # Connecting Arrow (between steps)
        if i < len(steps) - 1:
            arr = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, x + step_w + Inches(0.02), Inches(3.0), Inches(0.14), Inches(0.2)
            )
            arr.fill.solid()
            arr.fill.fore_color.rgb = TEXT_MUTED
            arr.line.fill.background()

    # Sequence Summary Callout Box at Bottom
    seq_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.1), Inches(11.733), Inches(1.5)
    )
    seq_box.fill.solid()
    seq_box.fill.fore_color.rgb = RGBColor(241, 245, 249)
    seq_box.line.color.rgb = BORDER_COLOR
    seq_box.line.width = Pt(1)

    tf_seq = seq_box.text_frame
    tf_seq.word_wrap = True
    tf_seq.margin_left = Inches(0.25)
    tf_seq.margin_top = Inches(0.15)

    p_sqh = tf_seq.paragraphs[0]
    p_sqh.text = "Key Operational Highlights across Workflow Stages"
    p_sqh.font.size = Pt(11)
    p_sqh.font.bold = True
    p_sqh.font.color.rgb = ROYAL_BLUE
    p_sqh.space_after = Pt(4)

    sq_pts = [
        "Immediate Feedback: Resume parsing executes in milliseconds upon upload, populating candidate fields for verification.",
        "Deterministic Scoring: Candidate-to-job matching computes instantaneously using exact mathematical weights — zero black-box hallucination.",
        "Audit Trail Integrity: Every stage advancement records updatedBy, timestamp, and review notes in MongoDB stageHistory array.",
        "Candidate Engagement: Applicants track application status in real-time, receiving upcoming interview alerts with Google Meet links."
    ]
    add_bullet_points(tf_seq, sq_pts, font_size=8.5, space_after=2)


def build_slide_10(prs):
    """SLIDE 10 — AI-ASSISTED FUNCTIONALITY & ALGORITHM"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "AI-Assisted Functionality & Matching Algorithm", "Scoring Engine & NLP Implementation")
    add_footer(slide, 10)

    # Left Column: Mathematical Formula & Weighting Breakdown
    create_card(slide, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.15))
    tb_l = slide.shapes.add_textbox(Inches(0.95), Inches(1.65), Inches(5.4), Inches(4.85))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "1. Multi-Factor Weighted Scoring Formula"
    p_lh.font.size = Pt(12)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ROYAL_BLUE
    p_lh.space_after = Pt(6)

    # Formula Box shape inside card
    f_box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(2.1), Inches(5.3), Inches(0.65)
    )
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
    f_box.line.color.rgb = BLUE_DARK
    f_box.line.width = Pt(1)
    tf_fb = f_box.text_frame
    p_fb = tf_fb.paragraphs[0]
    p_fb.text = "Overall Fit = (0.50 × Skills) + (0.25 × Exp) + (0.15 × Edu) + (0.10 × TF-IDF)"
    p_fb.font.size = Pt(9.5)
    p_fb.font.bold = True
    p_fb.font.color.rgb = BLUE_DARK
    p_fb.alignment = PP_ALIGN.CENTER

    tb_lw = slide.shapes.add_textbox(Inches(0.95), Inches(2.85), Inches(5.4), Inches(3.6))
    tf_lw = tb_lw.text_frame
    tf_lw.word_wrap = True

    weight_points = [
        "Skills Fit (50% Weight): Evaluates ratio of matched required skills against candidate skills and resume text. Includes synonym alias resolution (e.g., K8s -> Kubernetes, React.js -> React) plus bonus for extra taxonomy skills.",
        "Experience Fit (25% Weight): Compares candidate's total years against required minimum. Full score (100) if years ≥ required; curved ratio scoring with floor if below requirement.",
        "Education Fit (15% Weight): Compares degree ranks: PhD (Rank 5) > Master (4) > Bachelor (3) > Associate (2) > High School (1). Scores 100 if candidate meets or exceeds requirement.",
        "Semantic Fit (10% Weight): Leverages Natural library TF-IDF to compare job description top keywords against candidate resume text to measure contextual domain relevance."
    ]
    add_bullet_points(tf_lw, weight_points, font_size=8.5, space_after=4, bold_first_n_words=1)

    # Right Column: NLP Pipeline & Classification Tiers
    create_card(slide, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.15))
    tb_r = slide.shapes.add_textbox(Inches(6.95), Inches(1.65), Inches(5.4), Inches(4.85))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "2. NLP Pipeline & Recommendation Tiers"
    p_rh.font.size = Pt(12)
    p_rh.font.bold = True
    p_rh.font.color.rgb = ROYAL_BLUE
    p_rh.space_after = Pt(6)

    nlp_points = [
        "Document Parsing: Low-level binary buffer extraction via pdf-parse (PDFs) and mammoth (DOCX), yielding clean raw text.",
        "Entity & Contact Extraction: Compromise NLP extracts named entities (names, organizations); regular expressions extract emails, phone numbers, and portfolio links.",
        "Curated Skills Taxonomy: 500+ tech taxonomy categorizing languages, web frameworks, cloud/DevOps, AI/ML, and system design.",
        "Automated Recommendation Tiers:\n  • Strong Match (≥ 80%): Top priority interview recommendation.\n  • Good Match (65–79%): High alignment; minor skill/exp gaps.\n  • Potential Match (50–64%): Partial overlap; transferable skills.\n  • Low Match (< 50%): Significant qualification gap.",
        "Decoupled LLM Architecture: Modular service contract (fileBuffer, mimeType) -> ExtractedProfileJSON allows seamless future swapping with OpenAI or Google Gemini APIs!"
    ]
    add_bullet_points(tf_r, nlp_points, font_size=8.5, space_after=5, bold_first_n_words=1)


def build_slide_11(prs):
    """SLIDE 11 — DATABASE DESIGN"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Database Design & MongoDB Collections", "Data Modeling, Relationships & Schema Design")
    add_footer(slide, 11)

    card_w = Inches(5.7)
    card_h = Inches(2.45)
    xs = [Inches(0.8), Inches(6.8)]
    ys = [Inches(1.5), Inches(4.15)]

    collections = [
        (
            "Users Collection (User.js)",
            ROYAL_BLUE,
            [
                "_id: ObjectId (Primary Key)",
                "name: String (Required, trimmed, max 100)",
                "email: String (Required, unique index, regex validated)",
                "password: String (Bcrypt salted hash, select: false)",
                "role: String (Enum: ['recruiter', 'candidate'])",
                "company, title, phone, avatar: String",
                "timestamps: createdAt, updatedAt"
            ]
        ),
        (
            "Jobs Collection (Job.js)",
            BLUE_DARK,
            [
                "_id: ObjectId (Primary Key)",
                "title, company, department, location, type: String",
                "description: String (Required job overview)",
                "requiredSkills: [String] (At least 1 required)",
                "optionalSkills: [String] (Supplementary skills)",
                "minExperience, maxExperience: Number (Years)",
                "educationRequirement: Enum (High School to PhD)",
                "recruiter: ObjectId (Ref: User), status: Enum"
            ]
        ),
        (
            "CandidateProfiles Collection (CandidateProfile.js)",
            EMERALD,
            [
                "_id: ObjectId (Primary Key), user: ObjectId (Ref: User)",
                "originalFileName, fileUrl, mimeType, fileSize",
                "parsedData.name, email, phone, location, summary",
                "parsedData.skills: [String], technicalSkills, softSkills",
                "parsedData.totalYearsExperience: Number, highestEducation",
                "parsedData.experience: [{ company, title, durationYears, bullets }]",
                "parsedData.education: [{ institution, degree, fieldOfStudy }]",
                "rawText: String, parsingMetadata: { confidenceScore, processingTimeMs }"
            ]
        ),
        (
            "Applications Collection (Application.js)",
            RGBColor(147, 51, 234),
            [
                "_id: ObjectId (PK), job: Ref(Job), candidate: Ref(User)",
                "candidateProfile: ObjectId (Ref: CandidateProfile)",
                "stage: Enum ['Applied', 'Screened', 'Shortlisted', 'Interview Scheduled', 'Selected', 'Rejected']",
                "stageHistory: [{ stage, updatedBy: Ref(User), updatedAt, note }]",
                "matchAnalysis: { overallScore, skillsScore, experienceScore, educationScore, semanticScore, matchedSkills, missingSkills, recommendation }",
                "interviewDetails: { scheduledAt, type, meetingLink, interviewer, notes }",
                "Compound Index: { job: 1, candidate: 1 } (Unique Constraint)"
            ]
        ),
    ]

    idx = 0
    for r in range(2):
        for c in range(2):
            col_name, color, fields = collections[idx]
            create_card(slide, xs[c], ys[r], card_w, card_h)

            tb = slide.shapes.add_textbox(xs[c] + Inches(0.18), ys[r] + Inches(0.12), card_w - Inches(0.36), card_h - Inches(0.24))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = col_name
            p.font.size = Pt(11.5)
            p.font.bold = True
            p.font.color.rgb = color
            p.space_after = Pt(4)

            for f_text in fields:
                pf = tf.add_paragraph()
                pf.text = "• " + f_text
                pf.font.size = Pt(8.5)
                pf.font.name = "Consolas" if ":" in f_text else "Segoe UI"
                pf.font.color.rgb = TEXT_BODY
                pf.space_after = Pt(1.5)

            idx += 1


def build_slide_12(prs):
    """SLIDE 12 — USER INTERFACE SHOWCASE (Recruiter & Candidate)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "User Interface Showcase — Recruiter & Candidate Views", "Application Screens & User Experience")
    add_footer(slide, 12)

    # 3 High-Impact Screenshots side-by-side
    screens = [
        ("screenshots/3_recruiter_dashboard.png", "1. Recruiter Dashboard", "Live recruitment KPIs, active job listings, total applicants, and real-time activity stream."),
        ("screenshots/6_candidate_ranking.png", "2. Candidate Ranking & Breakdown", "Ranked candidate cards sorted by compatibility score (e.g. 96%), matched/missing skill tags & modal breakdown."),
        ("screenshots/5_kanban_pipeline.png", "3. 6-Stage Kanban Pipeline", "Interactive hiring board (Applied to Selected/Rejected) with 1-click stage advancement & interview scheduling.")
    ]

    sc_w = Inches(3.7)
    sc_h = Inches(2.3)
    start_x = Inches(0.8)
    gap = Inches(4.0)

    for i, (img_path, title, caption) in enumerate(screens):
        x = start_x + (i * gap)

        # Image Container Card
        create_card(slide, x, Inches(1.5), sc_w, Inches(3.5))

        # Screenshot Image
        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, x + Inches(0.1), Inches(1.6), width=sc_w - Inches(0.2))

        # Title
        t_box = slide.shapes.add_textbox(x + Inches(0.1), Inches(3.95), sc_w - Inches(0.2), Inches(0.3))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = ROYAL_BLUE

        # Caption
        c_box = slide.shapes.add_textbox(x + Inches(0.1), Inches(4.25), sc_w - Inches(0.2), Inches(0.7))
        tf_c = c_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = caption
        p_c.font.size = Pt(8.5)
        p_c.font.color.rgb = TEXT_BODY

    # Bottom Row: Candidate Views (Candidate Dashboard & Resume Uploader)
    cand_screens = [
        ("screenshots/7_candidate_dashboard.png", "4. Candidate Portal & Tracker", "Live application stage tracker, upcoming interview alerts with direct Google Meet links, and personalized job matches."),
        ("screenshots/8_resume_upload.png", "5. Modular Resume Ingestion", "Multi-format drag-and-drop resume uploader (PDF/DOCX/TXT) with real-time extracted entity verification & skill editing.")
    ]

    c_sc_w = Inches(5.7)
    c_start_x = Inches(0.8)
    c_gap = Inches(6.0)

    for j, (img_path, title, caption) in enumerate(cand_screens):
        x = c_start_x + (j * c_gap)
        create_card(slide, x, Inches(5.15), c_sc_w, Inches(1.5))

        if os.path.exists(img_path):
            slide.shapes.add_picture(img_path, x + Inches(0.1), Inches(5.25), width=Inches(2.1))

        t_box = slide.shapes.add_textbox(x + Inches(2.3), Inches(5.25), c_sc_w - Inches(2.4), Inches(0.3))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(10)
        p_t.font.bold = True
        p_t.font.color.rgb = ROYAL_BLUE

        c_box = slide.shapes.add_textbox(x + Inches(2.3), Inches(5.55), c_sc_w - Inches(2.4), Inches(0.9))
        tf_c = c_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = caption
        p_c.font.size = Pt(8.5)
        p_c.font.color.rgb = TEXT_BODY


def build_slide_13(prs):
    """SLIDE 13 — IMPLEMENTATION AND TESTING"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Implementation Highlights & Automated Testing", "Code Quality & Verification Results")
    add_footer(slide, 13)

    card_w = Inches(5.7)
    card_h = Inches(5.15)

    # Left Column: Implementation Architecture & Security
    create_card(slide, Inches(0.8), Inches(1.5), card_w, card_h)
    tb_l = slide.shapes.add_textbox(Inches(0.95), Inches(1.65), card_w - Inches(0.3), card_h - Inches(0.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "1. Implementation Highlights"
    p_lh.font.size = Pt(12)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ROYAL_BLUE
    p_lh.space_after = Pt(8)

    imp_points = [
        "Modular Separation of Concerns: Clean decoupling between Express routing, business controllers (job, application, resume), and independent service engines (resumeParser.js, matchingEngine.js).",
        "Stateless JWT Security: Token-based authentication with role-based route authorization (RBAC) protecting recruiter-only endpoints (job creation, stage advancement, ranking).",
        "Robust File Upload Pipeline: Multer-based multipart streaming enforcing 10MB limits, MIME-type whitelisting (PDF, DOCX, TXT), and disk storage hygiene.",
        "Mongoose Schema Validation: Strict validation rules, unique indexes, and pre-save password encryption hooks using BcryptJS with cost factor 10.",
        "Reactive Frontend Architecture: React 19 Context API (AuthContext) managing global session state; Axios interceptors injecting Bearer tokens automatically."
    ]
    add_bullet_points(tf_l, imp_points, font_size=8.5, space_after=6, bold_first_n_words=1)

    # Right Column: Verified Automated Integration Testing Suite
    create_card(slide, Inches(6.8), Inches(1.5), card_w, card_h)
    tb_r = slide.shapes.add_textbox(Inches(6.95), Inches(1.65), card_w - Inches(0.3), card_h - Inches(0.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "2. Automated Test Suite (server/test/api.test.js)"
    p_rh.font.size = Pt(12)
    p_rh.font.bold = True
    p_rh.font.color.rgb = ROYAL_BLUE
    p_rh.space_after = Pt(4)

    p_rsub = tf_r.add_paragraph()
    p_rsub.text = "12/12 Automated Integration Tests Executed & Verified (100% Pass Rate):"
    p_rsub.font.size = Pt(9)
    p_rsub.font.bold = True
    p_rsub.font.color.rgb = EMERALD
    p_rsub.space_after = Pt(6)

    test_cases = [
        "Test 1: Health check endpoint verification (/api/health ➔ 200 OK)",
        "Test 2: Recruiter login & JWT Bearer token generation",
        "Test 3: Candidate authentication & profile authorization",
        "Test 4: Job listings retrieval & dynamic query filters",
        "Test 5: Recruiter job creation & payload validation",
        "Test 6: Resume Parser NLP extraction (skills, contact, exp, degree)",
        "Test 7: Matching Engine 4-factor scoring calculation & recommendation",
        "Test 8: Recruiter ranked candidates endpoint retrieval",
        "Test 9: Kanban pipeline application fetching by job ID",
        "Test 10: Pipeline stage transition audit log & interview scheduling",
        "Test 11: Recruiter dashboard analytics metrics calculation",
        "Test 12: Candidate dashboard application tracking & test job cleanup"
    ]
    add_bullet_points(tf_r, test_cases, font_size=8, space_after=2)


def build_slide_14(prs):
    """SLIDE 14 — CONCLUSION AND FUTURE ENHANCEMENTS"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, LIGHT_BG)
    add_header(slide, "Conclusion & Future Enhancements", "Project Summary & Strategic Roadmap")
    add_footer(slide, 14)

    card_w = Inches(5.7)
    card_h = Inches(5.15)

    # Left Column: Project Conclusion
    create_card(slide, Inches(0.8), Inches(1.5), card_w, card_h)
    tb_l = slide.shapes.add_textbox(Inches(0.95), Inches(1.65), card_w - Inches(0.3), card_h - Inches(0.3))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "1. Project Summary & Completed Work"
    p_lh.font.size = Pt(12)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ROYAL_BLUE
    p_lh.space_after = Pt(8)

    concl_points = [
        "Full-Stack Integration: Successfully engineered a cohesive, production-ready recruitment automation platform leveraging the full MERN stack (MongoDB, Express, React 19, Node.js).",
        "Automated Resume Parsing: Eliminated manual data entry through modular NLP ingestion supporting PDF, DOCX, and TXT resumes, extracting contacts, 500+ taxonomy skills, and degrees.",
        "Objective Candidate Evaluation: Replaced subjective resume screening with a transparent, explainable 4-factor mathematical scoring formula (Skills, Experience, Education, TF-IDF).",
        "Visual Pipeline Management: Implemented an interactive 6-stage Kanban board with transition audit histories, recruiter feedback notes, and integrated interview scheduling.",
        "Candidate Transparency: Delivered a dedicated applicant portal providing real-time status tracking and direct Google Meet video access, removing hiring opacity."
    ]
    add_bullet_points(tf_l, concl_points, font_size=8.5, space_after=6, bold_first_n_words=1)

    # Right Column: Realistic Future Enhancements
    create_card(slide, Inches(6.8), Inches(1.5), card_w, card_h)
    tb_r = slide.shapes.add_textbox(Inches(6.95), Inches(1.65), card_w - Inches(0.3), card_h - Inches(0.3))
    tf_r = tb_r.text_frame
    tf_r.word_wrap = True

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "2. Realistic Future Enhancements"
    p_rh.font.size = Pt(12)
    p_rh.font.bold = True
    p_rh.font.color.rgb = ROYAL_BLUE
    p_rh.space_after = Pt(8)

    future_points = [
        "Large Language Model (LLM) Integration: Plug in Google Gemini 1.5 or OpenAI GPT-4o APIs via the existing modular resumeParser contract to extract nuanced soft skills and analyze cover letters.",
        "Automated Notification Gateway: Integrate NodeMailer and Twilio for automated candidate email/SMS alerts whenever an application advances or an interview is scheduled.",
        "Real-Time Collaborative WebSockets: Incorporate Socket.io to enable multi-recruiter live synchronization on the Kanban board with drag-and-drop mechanics.",
        "In-Browser Video & Code Pairing: Embed WebRTC video rooms and an interactive Monaco code editor for live technical interview screenings directly in the portal.",
        "Algorithmic Bias Auditing: Implement statistical fairness metrics to detect and prevent demographic skew in candidate ranking, ensuring equal opportunity hiring."
    ]
    add_bullet_points(tf_r, future_points, font_size=8.5, space_after=6, bold_first_n_words=1)


def build_slide_15(prs):
    """SLIDE 15 — THANK YOU (Dark Navy Closing Theme)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide, DARK_NAVY)

    # Decorative accent bar
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.12), Inches(5.8)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = CYAN_ACCENT
    accent_bar.line.fill.background()

    # Thank You Title
    t_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(7.5), Inches(1.3))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "THANK YOU!"
    p_t.font.size = Pt(44)
    p_t.font.bold = True
    p_t.font.color.rgb = WHITE
    p_t.font.name = "Segoe UI"

    # Subtitle / Questions
    q_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(7.5), Inches(0.8))
    tf_q = q_box.text_frame
    tf_q.word_wrap = True
    p_q = tf_q.paragraphs[0]
    p_q.text = "Questions & Discussion"
    p_q.font.size = Pt(22)
    p_q.font.color.rgb = CYAN_LIGHT
    p_q.font.name = "Segoe UI"

    # Tagline / Quote
    tag_box = slide.shapes.add_textbox(Inches(1.2), Inches(3.2), Inches(7.2), Inches(1.2))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = '"Empowering Talent Acquisition Teams with Explainable Algorithmic Intelligence, Modular NLP, and Transparent Hiring Pipelines."'
    p_tag.font.size = Pt(13)
    p_tag.font.italic = True
    p_tag.font.color.rgb = RGBColor(203, 213, 225)
    p_tag.font.name = "Segoe UI"

    p_foot = tf_tag.add_paragraph()
    p_foot.space_before = Pt(14)
    r_foot = p_foot.add_run()
    r_foot.text = "Project: AI-Assisted Recruitment System  |  FSD Mini Project"
    r_foot.font.size = Pt(11)
    r_foot.font.bold = True
    r_foot.font.color.rgb = CYAN_LIGHT
    r_foot.font.name = "Segoe UI"

    # Student Credentials Card (Right side)
    cred_card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.9), Inches(1.2), Inches(3.6), Inches(5.1)
    )
    cred_card.fill.solid()
    cred_card.fill.fore_color.rgb = CARD_DARK
    cred_card.line.color.rgb = BORDER_DARK
    cred_card.line.width = Pt(1.5)

    tf_cred = cred_card.text_frame
    tf_cred.word_wrap = True
    tf_cred.margin_left = Inches(0.25)
    tf_cred.margin_right = Inches(0.25)
    tf_cred.margin_top = Inches(0.25)

    p_ch = tf_cred.paragraphs[0]
    p_ch.text = "ACADEMIC DETAILS"
    p_ch.font.size = Pt(10.5)
    p_ch.font.bold = True
    p_ch.font.color.rgb = CYAN_LIGHT
    p_ch.font.name = "Segoe UI"
    p_ch.space_after = Pt(14)

    fields = [
        ("Presenter", "S. Srisakthi"),
        ("Register Number", "310624205269"),
        ("Degree & Branch", "B.Tech. Information Technology"),
        ("Department", "Information Technology (IT)"),
        ("Institution", "Easwari Engineering College\n(Autonomous), Chennai"),
        ("Project Type", "Full Stack Development Mini Project"),
        ("Date", "October 2026"),
    ]

    for label, val in fields:
        p_lbl = tf_cred.add_paragraph()
        p_lbl.text = label.upper()
        p_lbl.font.size = Pt(8)
        p_lbl.font.color.rgb = TEXT_MUTED
        p_lbl.font.name = "Segoe UI"
        p_lbl.space_after = Pt(1)

        p_val = tf_cred.add_paragraph()
        p_val.text = val
        p_val.font.size = Pt(10.5)
        p_val.font.bold = True
        p_val.font.color.rgb = WHITE
        p_val.font.name = "Segoe UI"
        p_val.space_after = Pt(8)


def main():
    print(f"Creating PowerPoint Presentation: {OUTPUT_FILE}...")
    prs = init_presentation()

    build_slide_1(prs)   # Slide 1: Title
    build_slide_2(prs)   # Slide 2: Introduction
    build_slide_3(prs)   # Slide 3: Problem Statement
    build_slide_4(prs)   # Slide 4: Project Objectives
    build_slide_5(prs)   # Slide 5: Existing vs Proposed System
    build_slide_6(prs)   # Slide 6: System Architecture
    build_slide_7(prs)   # Slide 7: Technologies Used
    build_slide_8(prs)   # Slide 8: System Modules
    build_slide_9(prs)   # Slide 9: Workflow
    build_slide_10(prs)  # Slide 10: AI-Assisted Functionality & Matching Algorithm
    build_slide_11(prs)  # Slide 11: Database Design
    build_slide_12(prs)  # Slide 12: User Interface Showcase
    build_slide_13(prs)  # Slide 13: Implementation & Testing
    build_slide_14(prs)  # Slide 14: Conclusion & Future Enhancements
    build_slide_15(prs)  # Slide 15: Thank You & Q&A

    prs.save(OUTPUT_FILE)
    print(f"[OK] Presentation saved successfully with {len(prs.slides)} slides to: {os.path.abspath(OUTPUT_FILE)}")


if __name__ == "__main__":
    main()
